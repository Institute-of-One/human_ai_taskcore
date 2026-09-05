"""Refuse to call the manuscript ready while anything in it is still unset.

The failure this exists to stop is specific and has happened: a document built
before a release was cut went to an editor still carrying a placeholder, because
the placeholder looked enough like ordinary text that nobody reading the prose
stopped at it. Nothing about the manuscript says which build it is, so the check
has to be a step rather than a habit.

Usage::

    python tools/presubmission_check.py
    python tools/presubmission_check.py --references   # also resolve every DOI

Exit status is non-zero if the manuscript is not ready to send.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
MANUSCRIPT = REPO / "paper" / "manuscript.md"
TEMPLATE = REPO / "paper" / "manuscript_template.md"
RELEASE = REPO / "results" / "release.json"

UNSET_MARKER = re.compile(r"\[UNSET: (\w+)[^\]]*\]")
PLACEHOLDER = re.compile(r"\{\{([a-z0-9_]+)\}\}")
TODO = re.compile(r"\[TODO[^\]]*\]")

#: Fields that must be filled before the manuscript can be sent anywhere.
REQUIRED_RELEASE_FIELDS = (
    "version_tag",
    "release_commit",
    "zenodo_version_doi",
)


def _git(*args):
    return subprocess.run(
        ["git", *args], cwd=REPO, capture_output=True, text=True
    )


#: The upload files Medical Physics asks for, and where this repository builds them.
BUILD = MANUSCRIPT.parent / "build"
DOCX = BUILD / "manuscript.docx"
TITLE_PAGE = BUILD / "title_page.docx"

#: Every one of these is a requirement a companion submission was returned for, or
#: one the submission form refused to advance without. None is hypothetical.
ABSTRACT_HEADINGS = ("Background", "Purpose", "Methods", "Results", "Conclusions")
ABSTRACT_WORD_LIMIT = 500
CANONICAL_AFFILIATION = "Institute of One, LISIT Co., Ltd., Tokyo 150-0044, Japan"


def _check_medical_physics(text: str) -> list[str]:
    """The format requirements, checked on the artefacts the editor receives.

    Medical Physics returned a companion submission before peer review for want of
    line numbers and page numbers, and its Files step will not advance without a
    separate title page. Page furniture exists only in the .docx, so a scan of the
    Markdown cannot see it: these read the built files.
    """
    import re  # noqa: PLC0415
    import zipfile  # noqa: PLC0415

    problems: list[str] = []

    abstract = text.partition("# Abstract")[2].partition("# 1.")[0]
    if not abstract.strip():
        problems.append("the manuscript has no Abstract section")
    else:
        for heading in ABSTRACT_HEADINGS:
            if f"**{heading}.**" not in abstract:
                problems.append(
                    f"the structured abstract has no {heading} heading; Medical "
                    f"Physics requires {', '.join(ABSTRACT_HEADINGS)}"
                )
        words = len(re.findall(r"[\w-]+", abstract))
        if words > ABSTRACT_WORD_LIMIT:
            problems.append(
                f"the abstract is {words} words, over the {ABSTRACT_WORD_LIMIT} allowed"
            )

    title = re.search(r'^title:\s*"(.+)"\s*$', text, re.M)
    if title is None:
        problems.append("the manuscript has no title in its front matter")

    if not DOCX.exists():
        problems.append("paper/build/manuscript.docx has not been built")
    else:
        with zipfile.ZipFile(DOCX) as archive:
            document = archive.read("word/document.xml").decode("utf-8")
            names = archive.namelist()
        if "lnNumType" not in document:
            problems.append(
                "the .docx carries no line numbering; Medical Physics returned a "
                "companion submission before peer review for exactly that"
            )
        if not any(name.startswith("word/footer") for name in names):
            problems.append("the .docx carries no footer, so no page numbers")
        stripped = re.sub(r"<[^>]+>", "", document)
        if "[@" in stripped:
            problems.append("an unresolved [@citation] key survives in the .docx")
        if "$" in stripped:
            problems.append("a dollar sign survives in the .docx: an equation printed as source")
        if "{{" in stripped:
            problems.append("an unresolved {{placeholder}} survives in the .docx")

    if not TITLE_PAGE.exists():
        problems.append(
            "paper/build/title_page.docx has not been built; Medical Physics will not "
            "advance the Files step without a separate title page"
        )
    elif title is not None:
        from docx import Document  # noqa: PLC0415

        page = " ".join(p.text for p in Document(str(TITLE_PAGE)).paragraphs)
        if title.group(1) not in page:
            problems.append("the title page and the manuscript carry different titles")
        if CANONICAL_AFFILIATION not in page:
            problems.append(
                f"the title page does not carry the canonical affiliation: "
                f"{CANONICAL_AFFILIATION}"
            )

    return problems


def check() -> list[str]:
    problems: list[str] = []
    text = MANUSCRIPT.read_text(encoding="utf-8")

    for field in sorted(set(UNSET_MARKER.findall(text))):
        problems.append(
            f"the manuscript still carries an unset release field: {field}"
        )
    if PLACEHOLDER.search(text):
        problems.append(
            "the rendered manuscript still has {{placeholders}}: re-run "
            "paper/make_figures.py"
        )
    for marker in sorted(set(TODO.findall(TEMPLATE.read_text(encoding="utf-8")))):
        problems.append(f"the template still carries {marker}")

    release = json.loads(RELEASE.read_text(encoding="utf-8"))
    for field in REQUIRED_RELEASE_FIELDS:
        if release.get(field) is None:
            problems.append(f"results/release.json: {field} is not set")

    # The release commit has to be a commit that exists, and the tag has to point
    # at it. A tag typed into the file and never cut is the same defect as an
    # unset field, only harder to see.
    commit = release.get("release_commit")
    tag = release.get("version_tag")
    if commit:
        if _git("cat-file", "-e", f"{commit}^{{commit}}").returncode:
            problems.append(f"release_commit {commit[:7]} is not a commit in this repository")
        elif tag:
            resolved = _git("rev-list", "-n", "1", tag)
            if resolved.returncode:
                problems.append(f"tag {tag} does not exist in this repository")
            elif resolved.stdout.strip() != commit:
                problems.append(
                    f"tag {tag} points at {resolved.stdout.strip()[:7]}, "
                    f"not at release_commit {commit[:7]}"
                )

    problems.extend(_check_medical_physics(text))

    # A release cut from a dirty tree archives something no commit describes.
    dirty = _git("status", "--porcelain").stdout.strip()
    if dirty:
        problems.append(
            f"the working tree has uncommitted changes ({len(dirty.splitlines())} paths)"
        )

    return problems


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument(
        "--references",
        action="store_true",
        help="also resolve every DOI in paper/references.bib against doi.org",
    )
    args = parser.parse_args(argv)

    problems = check()

    if args.references:
        sys.path.insert(0, str(Path(__file__).parent))
        import check_references

        if check_references.main([]):
            problems.append("paper/references.bib has entries that do not check out")

    if problems:
        print("NOT READY TO SUBMIT:\n")
        for problem in problems:
            print(f"  - {problem}")
        print(
            "\nThe release fields are filled after the GitHub release is cut and "
            "Zenodo has minted the version DOI. Re-run paper/make_figures.py "
            "afterwards so the manuscript picks them up."
        )
        return 1

    print("ready: no unset fields, no placeholders, no TODOs, tree is clean")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
